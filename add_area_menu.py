#!/usr/bin/env python3
"""
Add Area and Menu Items columns to a leads xlsx using Gemini + Google Search grounding.

Usage:
    py add_area_menu.py "Downtown Detroit leads.xlsx"

Output overwrites the input file with columns reordered to:
    Name | Area | Menu Items | Website | Email | Instagram | Facebook
"""

import os, re, sys, json, time, requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from dotenv import load_dotenv

load_dotenv()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
if not GEMINI_API_KEY:
    print("ERROR: GEMINI_API_KEY not in .env")
    sys.exit(1)

GEMINI_URL = (
    "https://generativelanguage.googleapis.com/v1beta/models/"
    "gemini-2.5-flash:generateContent?key=" + GEMINI_API_KEY
)

RATE_DELAY = 4  # seconds between calls


PROMPT = """Use Google Search to look up this restaurant and return ONLY this JSON (no markdown, no explanation):

Restaurant: {name}
City: {city}
Website: {website}

{{
  "area": "short neighborhood name only, e.g. Greektown or Midtown (NOT a sentence, max 4 words, null if unknown)",
  "menu_items": "up to 5 popular dish or drink names comma-separated, empty string if unknown"
}}"""


def call_gemini_search(prompt: str) -> str:
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "tools": [{"google_search": {}}],
        "generationConfig": {"temperature": 0.2, "maxOutputTokens": 2048},
    }
    r = requests.post(GEMINI_URL, json=payload, timeout=60)
    r.raise_for_status()
    data = r.json()
    parts = data["candidates"][0].get("content", {}).get("parts", [])
    text_parts = [p["text"] for p in parts if "text" in p]
    if not text_parts:
        raise ValueError(f"No text in response")
    return text_parts[-1].strip()


def parse_json(text: str) -> dict:
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    m = re.search(r"\{[\s\S]*\}", text)
    if m:
        text = m.group(0)
    return json.loads(text)


def get_area_menu(name: str, website: str, city: str) -> dict:
    prompt = PROMPT.format(name=name, city=city, website=website or "unknown")
    try:
        raw = call_gemini_search(prompt)
        return parse_json(raw)
    except Exception as e:
        print(f"    ERROR: {e}")
        return {"area": None, "menu_items": ""}


def main():
    if len(sys.argv) < 2:
        print("Usage: py add_area_menu.py <input.xlsx> [city]")
        sys.exit(1)

    input_path = sys.argv[1]
    city = sys.argv[2] if len(sys.argv) > 2 else "Detroit, Michigan"

    if not os.path.exists(input_path):
        print(f"ERROR: File not found: {input_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(input_path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    header = list(rows[0])
    data_rows = [list(r) for r in rows[1:] if r[0]]

    col = {str(h).strip(): i for i, h in enumerate(header) if h}
    name_i    = col.get("Name", 0)
    website_i = col.get("Website")
    email_i   = col.get("Email")
    ig_i      = col.get("Instagram")
    fb_i      = col.get("Facebook")
    area_i    = col.get("Area")
    menu_i    = col.get("Menu Items")

    OUT_HEADER = ["Name", "Area", "Menu Items", "Website", "Email", "Instagram", "Facebook"]

    total = len(data_rows)
    results = []

    for idx, row in enumerate(data_rows, 1):
        name    = row[name_i] or ""
        website = row[website_i] if website_i is not None else ""
        email   = row[email_i]   if email_i   is not None else ""
        ig      = row[ig_i]      if ig_i      is not None else ""
        fb      = row[fb_i]      if fb_i      is not None else ""

        # If area/menu already present in source, carry them over
        existing_area = (row[area_i] if area_i is not None else None) or ""
        existing_menu = (row[menu_i] if menu_i is not None else None) or ""

        if existing_area and existing_menu:
            print(f"[{idx}/{total}] SKIP (already filled): {name}")
            results.append([name, existing_area, existing_menu, website, email, ig, fb])
            continue

        print(f"[{idx}/{total}] Looking up: {name}")
        info = get_area_menu(name, website, city)
        area       = info.get("area") or existing_area or ""
        menu_items = info.get("menu_items") or existing_menu or ""

        print(f"    Area       : {area or '-'}")
        print(f"    Menu Items : {menu_items or '-'}")
        results.append([name, area, menu_items, website, email, ig, fb])

        if idx < total:
            time.sleep(RATE_DELAY)

    # Write output back to same file
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Leads"

    ws_out.append(OUT_HEADER)
    for cell in ws_out[1]:
        cell.font      = Font(bold=True, color="FFFFFF")
        cell.fill      = PatternFill(fill_type="solid", fgColor="2E4057")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, row in enumerate(results):
        ws_out.append(row)
        bg = "FFFFFF" if i % 2 == 0 else "EDF2F7"
        for c in range(1, len(OUT_HEADER) + 1):
            cell = ws_out.cell(i + 2, c)
            cell.fill      = PatternFill(fill_type="solid", fgColor=bg)
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    # Column widths
    for col_letter, width in zip("ABCDEFG", [38, 22, 55, 42, 38, 42, 42]):
        ws_out.column_dimensions[col_letter].width = width

    wb_out.save(input_path)
    print(f"\nDone. {len(results)} rows saved to {input_path}")


if __name__ == "__main__":
    main()
